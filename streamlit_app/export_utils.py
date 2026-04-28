"""PDF, Word (.docx), and Excel export — ported from src/utils/exportUtils.ts."""

from __future__ import annotations

import io
import re
from datetime import date
from typing import Any, Callable, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

from streamlit_app.api_service import (
    ACADEMIC_CORPORATE_SUBMETRIC_IDS,
    COLLABORATION_SUBMETRIC_IDS,
)

MetricGetter = Callable[[Dict[str, Any]], Dict[str, Any]]


def _sanitize_filename(filename: str) -> str:
    return (
        re.sub(r'[<>:"/\\|?*]', "", filename)
        .replace("..", "")
        .lstrip(".-")[:100]
    )


def _sanitize_text(text: Optional[str], limit: int = 500) -> str:
    if not text:
        return ""
    return (
        text.replace("<", "")
        .replace(">", "")
        .replace("javascript:", "")
        .replace("data:", "")[:limit]
    )


def _validate_export_data(data: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not isinstance(data, list):
        raise ValueError("Invalid export data: must be a list")
    if len(data) > 100:
        raise ValueError("Too many authors for export (max 100)")
    out = []
    for item in data:
        if not isinstance(item, dict):
            raise ValueError("Invalid export data item")
        aid = item.get("authorId")
        if not aid or not isinstance(aid, str):
            raise ValueError("Invalid author ID")
        sm = item.get("selectedMetrics")
        mo = item.get("metricOrder")
        out.append(
            {
                **item,
                "authorId": _sanitize_text(aid),
                "authorName": _sanitize_text(item.get("authorName")),
                "selectedMetrics": (
                    [m for m in sm if isinstance(m, str)][:20]
                    if isinstance(sm, list)
                    else None
                ),
                "metricOrder": (
                    [m for m in mo if isinstance(m, str)][:20]
                    if isinstance(mo, list)
                    else None
                ),
            }
        )
    return out


METRIC_DEFS: List[Dict[str, Any]] = [
    {
        "id": "publication",
        "label": "Publication",
        "getData": lambda m: m["scholarlyOutput"],
        "isYearBased": True,
    },
    {
        "id": "fwci",
        "label": "FWCI",
        "getData": lambda m: m["fwci"],
        "isYearBased": True,
    },
    {
        "id": "topJournal",
        "label": "Publications in Top 10% Journals",
        "getData": lambda m: m["topJournal"],
        "isYearBased": True,
        "suffix": "%",
    },
    {
        "id": "citationCount",
        "label": "Citation Count",
        "getData": lambda m: m["citationCount"],
        "isYearBased": True,
    },
    {
        "id": "citationsPerPublication",
        "label": "Citations Per Publication",
        "getData": lambda m: m["citationsPerPublication"],
        "isYearBased": True,
    },
    {
        "id": "hIndex",
        "label": "H-Index",
        "getData": lambda m: {"byYear": {}, "total": m["hIndex"]["value"]},
        "isYearBased": False,
    },
    {
        "id": "collaborationInternational",
        "label": "International collaboration %",
        "getData": lambda m: m.get("collaborationInternational")
        or {"byYear": {}, "total": "N/A"},
        "isYearBased": True,
        "isCollaboration": True,
        "suffix": "%",
    },
    {
        "id": "collaborationNational",
        "label": "National collaboration %",
        "getData": lambda m: m.get("collaborationNational")
        or {"byYear": {}, "total": "N/A"},
        "isYearBased": True,
        "isCollaboration": True,
        "suffix": "%",
    },
    {
        "id": "collaborationInstitutional",
        "label": "Institutional collaboration %",
        "getData": lambda m: m.get("collaborationInstitutional")
        or {"byYear": {}, "total": "N/A"},
        "isYearBased": True,
        "isCollaboration": True,
        "suffix": "%",
    },
    {
        "id": "collaborationSingleAuthorship",
        "label": "Single authorship %",
        "getData": lambda m: m.get("collaborationSingleAuthorship")
        or {"byYear": {}, "total": "N/A"},
        "isYearBased": True,
        "isCollaboration": True,
        "suffix": "%",
    },
    {
        "id": "academicCorporateWith",
        "label": "Academic–corporate collaboration %",
        "getData": lambda m: m.get("academicCorporateWith")
        or {"byYear": {}, "total": "N/A"},
        "isYearBased": True,
        "isCollaboration": True,
        "suffix": "%",
    },
    {
        "id": "academicCorporateWithout",
        "label": "No academic–corporate collaboration %",
        "getData": lambda m: m.get("academicCorporateWithout")
        or {"byYear": {}, "total": "N/A"},
        "isYearBased": True,
        "isCollaboration": True,
        "suffix": "%",
    },
]


def _get_selected_metrics(author_data: Dict[str, Any]) -> List[Dict[str, Any]]:
    selected_ids = author_data.get("selectedMetrics") or [m["id"] for m in METRIC_DEFS]
    order = author_data.get("metricOrder") or [m["id"] for m in METRIC_DEFS]
    by_id = {m["id"]: m for m in METRIC_DEFS}
    result = []
    for mid in order:
        if mid in selected_ids and mid in by_id:
            result.append(by_id[mid])
    return result


def _get_dynamic_years(author_data: Dict[str, Any]) -> List[str]:
    ds = author_data.get("dataSource") or {}
    ms = ds.get("metricStartYear")
    me = ds.get("metricEndYear")
    if ms is not None and me is not None:
        return [str(y) for y in range(int(ms), int(me) + 1)]
    metrics = author_data.get("metrics") or {}
    keys = [
        metrics.get("scholarlyOutput", {}).get("byYear") or {},
        metrics.get("fwci", {}).get("byYear") or {},
        metrics.get("topJournal", {}).get("byYear") or {},
        metrics.get("citationCount", {}).get("byYear") or {},
        metrics.get("citationsPerPublication", {}).get("byYear") or {},
        *[
            (metrics.get(cid) or {}).get("byYear") or {}
            for cid in COLLABORATION_SUBMETRIC_IDS
        ],
        *[
            (metrics.get(aid) or {}).get("byYear") or {}
            for aid in ACADEMIC_CORPORATE_SUBMETRIC_IDS
        ],
    ]
    years: set[int] = set()
    for yd in keys:
        for y in yd:
            try:
                years.add(int(y))
            except (TypeError, ValueError):
                pass
    if years:
        return [str(y) for y in sorted(years)]
    return ["2019", "2020", "2021", "2022", "2023", "2024"]


def _format_export_value(
    value: Any, suffix: str = "", metric_id: Optional[str] = None
) -> str:
    if value == "N/A" or value is None:
        return "N/A"
    if isinstance(value, (int, float)):
        if suffix == "%":
            fv = float(value)
            return f"{int(fv)}%" if fv.is_integer() else f"{fv:.2f}%"
        if metric_id in ("fwci", "citationsPerPublication", "topJournal"):
            return f"{float(value):.2f}{suffix}"
        return f"{round(value)}{suffix}"
    return f"{value}{suffix}"


def _cell_year(
    by_year: Dict[str, float],
    year: str,
    is_collab: bool,
    metric_id: Optional[str],
) -> str:
    v = by_year.get(year)
    if v is None:
        return "N/A"
    if is_collab or metric_id == "topJournal":
        return f"{int(v)}%" if v % 1 == 0 else f"{float(v):.2f}%"
    if metric_id in ("fwci", "citationsPerPublication"):
        return f"{float(v):.2f}"
    return str(round(v))


def _pdf_safe(s: str) -> str:
    return "".join(c if 32 <= ord(c) < 127 or c in "\n\t" else "?" for c in (s or ""))


def _pdf_wrap_width(pdf: FPDF) -> float:
    """Usable width for ``multi_cell`` — never pass ``w=0`` (remaining width can be 0 if ``x`` is wrong)."""
    try:
        epw = getattr(pdf, "epw", None)
        if epw is not None and float(epw) > 0:
            return max(float(epw), 20.0)
    except (TypeError, ValueError):
        pass
    w = float(pdf.w) - float(pdf.l_margin) - float(pdf.r_margin)
    return max(w, 20.0)


def _pdf_reset_x_margin(pdf: FPDF) -> None:
    pdf.set_x(pdf.l_margin)


def export_pdf_bytes(data: List[Dict[str, Any]], filename_base: str = "research-metrics") -> Tuple[bytes, str]:
    validated = _validate_export_data(data)
    name = _sanitize_filename(filename_base)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, _pdf_safe("SciVal Research Metrics Report"), ln=True, align="C")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 8, _pdf_safe(f"Generated on: {date.today().isoformat()}"), ln=True, align="C")
    pdf.ln(6)

    for idx, author_data in enumerate(validated):
        _pdf_reset_x_margin(pdf)
        if pdf.get_y() > 250:
            pdf.add_page()
            _pdf_reset_x_margin(pdf)

        title = (
            f"{author_data.get('authorName') or ''} (ID: {author_data['authorId']})"
            if author_data.get("authorName")
            else f"Author ID: {author_data['authorId']}"
        )
        wrap_w = _pdf_wrap_width(pdf)
        pdf.set_font("Helvetica", "B", 12)
        _pdf_reset_x_margin(pdf)
        pdf.multi_cell(wrap_w, 8, _pdf_safe(title))
        _pdf_reset_x_margin(pdf)
        ds = author_data.get("dataSource")
        if ds:
            wrap_w = _pdf_wrap_width(pdf)
            # Single line "Source: …" (matches Excel/Word); multi_cell still wraps if very long.
            pdf.set_font("Helvetica", "", 9)
            _pdf_reset_x_margin(pdf)
            sn = str(ds.get("sourceName", "") or "—")
            pdf.multi_cell(wrap_w, 5, _pdf_safe(f"Source: {sn}"))
            _pdf_reset_x_margin(pdf)
            pdf.multi_cell(
                wrap_w,
                5,
                _pdf_safe(f"Last updated: {ds.get('lastUpdated', '')}"),
            )
            _pdf_reset_x_margin(pdf)
            ms, me = ds.get("metricStartYear", ""), ds.get("metricEndYear", "")
            if ms or me:
                pdf.multi_cell(
                    wrap_w,
                    5,
                    _pdf_safe(f"Metric period: {ms} - {me}"),
                )
                _pdf_reset_x_margin(pdf)
            pdf.ln(2)

        selected = _get_selected_metrics(author_data)
        years = _get_dynamic_years(author_data)
        metrics = author_data.get("metrics") or {}

        if not selected:
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 8, _pdf_safe("No metrics selected for export"), ln=True)
            pdf.ln(4)
            continue

        col_w = [40] + [18] * len(years) + [22]
        total_w = sum(col_w)
        if total_w > 190:
            scale = 190 / total_w
            col_w = [int(c * scale) for c in col_w]

        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(59, 130, 246)
        pdf.set_text_color(255, 255, 255)
        headers = ["Metric"] + years + ["Total / Avg"]
        for i, h in enumerate(headers):
            pdf.cell(col_w[i], 7, _pdf_safe(str(h)), border=1, fill=True)
        pdf.ln()
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 7)

        fill = False
        for metric in selected:
            get_data: MetricGetter = metric["getData"]
            d = get_data(metrics)
            is_collab = bool(metric.get("isCollaboration"))
            mid = metric["id"]
            suffix = metric.get("suffix") or ""

            if metric["isYearBased"]:
                row = [metric["label"]] + [
                    _cell_year(d.get("byYear") or {}, y, is_collab, mid) for y in years
                ] + [_format_export_value(d.get("total"), suffix, mid)]
            else:
                row = [metric["label"]] + ["N/A"] * len(years) + [
                    _format_export_value(d.get("total"), suffix, mid)
                ]

            if fill:
                pdf.set_fill_color(248, 250, 252)
            else:
                pdf.set_fill_color(255, 255, 255)
            for i, cell in enumerate(row):
                pdf.cell(
                    col_w[i],
                    6,
                    _pdf_safe(str(cell)),
                    border=1,
                    fill=True,
                )
            pdf.ln()
            fill = not fill

        if idx < len(validated) - 1:
            pdf.ln(6)
            pdf.set_draw_color(200, 200, 200)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())

    raw = pdf.output(dest="S")
    if isinstance(raw, str):
        data = raw.encode("latin-1", errors="replace")
    else:
        # fpdf2 may return bytearray; Streamlit download_button requires bytes
        data = bytes(raw)
    return data, f"{name}.pdf"


def export_excel_bytes(data: List[Dict[str, Any]], filename_base: str = "research-metrics") -> Tuple[bytes, str]:
    validated = _validate_export_data(data)
    name = _sanitize_filename(filename_base)
    wb = Workbook()
    wb.remove(wb.active)

    for index, author_data in enumerate(validated):
        ws = wb.create_sheet(
            title=f"Author_{index + 1}" if len(validated) > 1 else "Metrics"
        )
        row = 1
        ws.cell(row, 1, "SciVal Research Metrics Report").font = Font(bold=True, size=14)
        row += 1
        line2 = (
            f"Author: {author_data.get('authorName', '')} (ID: {author_data['authorId']})"
            if author_data.get("authorName")
            else f"Author ID: {author_data['authorId']}"
        )
        ws.cell(row, 1, line2)
        row += 1
        ws.cell(row, 1, f"Generated on: {date.today().isoformat()}")
        row += 2

        ds = author_data.get("dataSource")
        if ds:
            ws.cell(row, 1, f"Source: {ds.get('sourceName', '')}")
            row += 1
            ws.cell(row, 1, f"Last Updated: {ds.get('lastUpdated', '')}")
            row += 1
            ws.cell(
                row,
                1,
                f"Period: {ds.get('metricStartYear', '')} - {ds.get('metricEndYear', '')}",
            )
            row += 2

        years = _get_dynamic_years(author_data)
        selected = _get_selected_metrics(author_data)
        metrics = author_data.get("metrics") or {}

        if not selected:
            ws.cell(row, 1, "No metrics selected for export")
            continue

        headers = ["Metric"] + years + ["Total / Avg"]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row, c, h)
            cell.font = Font(bold=True)
        row += 1

        for metric in selected:
            get_data: MetricGetter = metric["getData"]
            d = get_data(metrics)
            mid = metric["id"]
            if metric["isYearBased"]:
                cells = [metric["label"]]
                for y in years:
                    v = (d.get("byYear") or {}).get(y)
                    if v is not None:
                        if metric.get("isCollaboration") or mid == "topJournal":
                            cells.append(
                                f"{int(v)}%"
                                if isinstance(v, (int, float)) and v % 1 == 0
                                else f"{float(v):.2f}%"
                            )
                        elif mid in ("fwci", "citationsPerPublication"):
                            cells.append(f"{float(v):.2f}")
                        else:
                            cells.append(str(round(v)))
                    else:
                        cells.append("N/A")
                tot = d.get("total")
                if isinstance(tot, (int, float)):
                    if metric.get("isCollaboration") or mid == "topJournal":
                        cells.append(
                            f"{int(tot)}%"
                            if tot % 1 == 0
                            else f"{float(tot):.2f}%"
                        )
                    elif mid in ("fwci", "citationsPerPublication"):
                        cells.append(f"{float(tot):.2f}")
                    else:
                        cells.append(str(round(tot)))
                else:
                    cells.append(str(tot))
            else:
                tot = d.get("total")
                cells = [metric["label"]] + ["N/A"] * len(years)
                if isinstance(tot, (int, float)):
                    if mid in ("fwci", "citationsPerPublication"):
                        cells.append(f"{float(tot):.2f}")
                    elif mid == "topJournal":
                        cells.append(
                            f"{int(tot)}%"
                            if tot % 1 == 0
                            else f"{float(tot):.2f}%"
                        )
                    else:
                        cells.append(str(round(tot)))
                else:
                    cells.append(str(tot))
            for c, val in enumerate(cells, 1):
                ws.cell(row, c, val)
            row += 1

        ws.column_dimensions["A"].width = 28
        for i in range(len(years)):
            col_letter = chr(ord("B") + i) if i < 26 else "Z"
            ws.column_dimensions[col_letter].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    # Ensure strict bytes for st.download_button
    return bytes(buf.getvalue()), f"{name}.xlsx"


def export_docx_bytes(data: List[Dict[str, Any]], filename_base: str = "research-metrics") -> Tuple[bytes, str]:
    """Export the same metric tables as Excel, as a Word (.docx) document."""
    validated = _validate_export_data(data)
    name = _sanitize_filename(filename_base)
    doc = Document()
    h0 = doc.add_heading("SciVal Research Metrics Report", 0)
    h0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0 = doc.add_paragraph(f"Generated on: {date.today().isoformat()}")
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    for index, author_data in enumerate(validated):
        if index > 0:
            doc.add_page_break()

        if len(validated) > 1:
            doc.add_heading(f"Author {index + 1}", level=1)

        line2 = (
            f"Author: {author_data.get('authorName', '')} (ID: {author_data['authorId']})"
            if author_data.get("authorName")
            else f"Author ID: {author_data['authorId']}"
        )
        doc.add_paragraph(line2)

        ds = author_data.get("dataSource")
        if ds:
            doc.add_paragraph(f"Source: {ds.get('sourceName', '')}")
            doc.add_paragraph(f"Last Updated: {ds.get('lastUpdated', '')}")
            doc.add_paragraph(
                f"Period: {ds.get('metricStartYear', '')} - {ds.get('metricEndYear', '')}"
            )

        years = _get_dynamic_years(author_data)
        selected = _get_selected_metrics(author_data)
        metrics = author_data.get("metrics") or {}

        if not selected:
            doc.add_paragraph("No metrics selected for export")
            continue

        headers = ["Metric"] + years + ["Total / Avg"]
        ncols = len(headers)
        table = doc.add_table(rows=1, cols=ncols)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        for c, h in enumerate(headers):
            hdr_cells[c].text = str(h)
            for run in hdr_cells[c].paragraphs[0].runs:
                run.bold = True

        for metric in selected:
            get_data: MetricGetter = metric["getData"]
            d = get_data(metrics)
            mid = metric["id"]
            row_cells = table.add_row().cells
            if metric["isYearBased"]:
                cells = [metric["label"]]
                for y in years:
                    v = (d.get("byYear") or {}).get(y)
                    if v is not None:
                        if metric.get("isCollaboration") or mid == "topJournal":
                            cells.append(
                                f"{int(v)}%"
                                if isinstance(v, (int, float)) and v % 1 == 0
                                else f"{float(v):.2f}%"
                            )
                        elif mid in ("fwci", "citationsPerPublication"):
                            cells.append(f"{float(v):.2f}")
                        else:
                            cells.append(str(round(v)))
                    else:
                        cells.append("N/A")
                tot = d.get("total")
                if isinstance(tot, (int, float)):
                    if metric.get("isCollaboration") or mid == "topJournal":
                        cells.append(
                            f"{int(tot)}%"
                            if tot % 1 == 0
                            else f"{float(tot):.2f}%"
                        )
                    elif mid in ("fwci", "citationsPerPublication"):
                        cells.append(f"{float(tot):.2f}")
                    else:
                        cells.append(str(round(tot)))
                else:
                    cells.append(str(tot))
            else:
                tot = d.get("total")
                cells = [metric["label"]] + ["N/A"] * len(years)
                if isinstance(tot, (int, float)):
                    if mid in ("fwci", "citationsPerPublication"):
                        cells.append(f"{float(tot):.2f}")
                    elif mid == "topJournal":
                        cells.append(
                            f"{int(tot)}%"
                            if tot % 1 == 0
                            else f"{float(tot):.2f}%"
                        )
                    else:
                        cells.append(str(round(tot)))
                else:
                    cells.append(str(tot))
            for c, val in enumerate(cells):
                row_cells[c].text = str(val)

        try:
            table.columns[0].width = Inches(1.9)
        except (AttributeError, ValueError, TypeError):
            pass

    buf = io.BytesIO()
    doc.save(buf)
    return bytes(buf.getvalue()), f"{name}.docx"
